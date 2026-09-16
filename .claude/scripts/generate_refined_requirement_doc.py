#!/usr/bin/env python3
"""
Generate a formal Word report from a JSON refinement package.

Usage:
  python .claude/scripts/generate_refined_requirement_doc.py \
      --input refinement_result.json \
      --output output/Refined_Requirement_<ID>.docx

The script intentionally contains formatting logic only. Requirement
reasoning, evidence classification, scoring, and refinement decisions
belong to the Requirements Refinement Agent and its SKILL.md.
"""

import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def safe_name(value):
    value = str(value or "REQUIREMENT")
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("_")
    return value[:100] or "REQUIREMENT"

def add_value(doc, heading, value):
    doc.add_heading(heading, level=2)
    if value is None or value == "":
        doc.add_paragraph("Not Applicable — no information was supplied or established.")
        return
    if isinstance(value, list):
        for item in value:
            if isinstance(item, dict):
                doc.add_paragraph(" | ".join(f"{k}: {v}" for k, v in item.items()))
            else:
                doc.add_paragraph(str(item), style="List Bullet")
    elif isinstance(value, dict):
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.rows[0].cells[0].text = "Attribute"
        table.rows[0].cells[1].text = "Value"
        for k, v in value.items():
            cells = table.add_row().cells
            cells[0].text = str(k)
            cells[1].text = str(v)
    else:
        doc.add_paragraph(str(value))

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    doc = Document()
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)

    title = doc.add_heading("REQUIREMENT REFINEMENT REPORT", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Generated: ").bold = True
    p.add_run(datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z"))

    doc.add_paragraph("Approval status: NOT APPROVED — awaiting BA/PO review.").bold = True

    ordered = [
        "requirement_id", "original_requirement", "refined_requirement",
        "business_objective", "structural_extraction", "actors_and_roles",
        "in_scope_behavior", "out_of_scope_behavior", "functional_behavior",
        "business_rules", "inputs_outputs", "validation_rules",
        "error_exception_handling", "positive_behavior", "negative_behavior",
        "boundary_edge_behavior", "dependencies", "integration_requirements",
        "non_functional_requirements", "security_authorization",
        "auditability_observability", "assumptions", "constraints",
        "missing_information_gaps", "ambiguities", "clarification_questions",
        "unknowns", "conflicts", "related_duplicate_requirements",
        "version_comparison", "change_impact_analysis", "testability_assessment",
        "automation_feasibility_assessment", "confirmed_acceptance_criteria",
        "proposed_acceptance_criteria", "traceability_candidates",
        "requirements_critic_findings", "validation_result", "quality_score",
        "refinement_status", "ba_po_review_summary"
    ]
    labels = {
        "requirement_id":"1. Requirement ID", "original_requirement":"2. Original Requirement",
        "refined_requirement":"3. Refined Requirement",
        "business_objective":"4. Business Objective",
        "structural_extraction":"5. Structural Extraction",
        "actors_and_roles":"6. Actors and Roles", "in_scope_behavior":"7. In-Scope Behavior",
        "out_of_scope_behavior":"8. Out-of-Scope Behavior", "functional_behavior":"9. Functional Behavior",
        "business_rules":"10. Business Rules", "inputs_outputs":"11. Inputs and Outputs",
        "validation_rules":"12. Validation Rules", "error_exception_handling":"13. Error and Exception Handling",
        "positive_behavior":"14. Positive Behavior", "negative_behavior":"15. Negative Behavior",
        "boundary_edge_behavior":"16. Boundary and Edge Behavior", "dependencies":"17. Dependencies",
        "integration_requirements":"18. Integration Requirements",
        "non_functional_requirements":"19. Non-Functional Requirements",
        "security_authorization":"20. Security and Authorization",
        "auditability_observability":"21. Auditability and Observability",
        "assumptions":"22. Assumptions", "constraints":"23. Constraints",
        "missing_information_gaps":"24. Missing Information / Gaps",
        "ambiguities":"25. Ambiguities", "clarification_questions":"26. Clarification Questions",
        "unknowns":"27. Unknowns", "conflicts":"28. Conflicts or Inconsistencies",
        "related_duplicate_requirements":"29. Related / Duplicate Requirements",
        "version_comparison":"30. Version Comparison",
        "change_impact_analysis":"31. Change Impact Analysis",
        "testability_assessment":"32. Testability Assessment",
        "automation_feasibility_assessment":"33. Automation Feasibility Assessment",
        "confirmed_acceptance_criteria":"34. Confirmed Acceptance Criteria",
        "proposed_acceptance_criteria":"35. Proposed Acceptance Criteria — Requires BA/PO Confirmation",
        "traceability_candidates":"36. Traceability Candidates",
        "requirements_critic_findings":"37. Requirements Critic Findings",
        "validation_result":"38. Validation Result", "quality_score":"39. Quality Score",
        "refinement_status":"40. Refinement Status",
        "ba_po_review_summary":"41. BA/PO Review Package Summary"
    }
    for key in ordered:
        if key in data:
            add_value(doc, labels[key], data[key])

    doc.add_paragraph("Approval status: NOT APPROVED — awaiting BA/PO review.").bold = True

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(str(output))

if __name__ == "__main__":
    main()
